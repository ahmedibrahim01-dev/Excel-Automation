import openpyxl

workbook = openpyxl.load_workbook("Organ.xlsx")

sheet = workbook.active

result_workbook = openpyxl.Workbook()

result_sheet = result_workbook.active

result_sheet["A1"] = "name"
result_sheet["B1"] = "total"

for row in sheet.iter_rows(min_row=5, values_only=True):

    if row[1] is None :
        continue

    name = row[1]
    price = row[2]
    quantity = row[3]
    total = price * quantity
    result_sheet.append([name,total])

result_workbook.save("results.xlsx")
